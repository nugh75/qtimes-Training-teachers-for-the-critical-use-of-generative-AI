#!/usr/bin/env python3
"""Build canonical corpus markdown files from immutable XLSX sources in dati/.

This script formalizes the first reproducible step of the PRAXIS pipeline:

    dati/ (*.xlsx) -> 02-corpus/*.md (legacy QA format)

The output format is intentionally compatible with `label_raw_with_ollama.py`,
which can parse these files with its legacy parser before restructuring.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from collections import Counter
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
except Exception as exc:  # pragma: no cover - import path depends on env
    raise SystemExit(
        "Missing dependency: openpyxl. Install it before running this script."
    ) from exc


ROOT_DIR = Path(__file__).resolve().parents[1] / "Praxis-ql"
REPO_ROOT = ROOT_DIR.parent

DEFAULT_DATI_DIR = REPO_ROOT / "dati"
DEFAULT_OUTPUT_DIR = ROOT_DIR / "02-corpus"
DEFAULT_MANIFEST_OUT = ROOT_DIR / "04-label" / "labels" / "corpus-build-manifest.json"
DEFAULT_OVERRIDES = ROOT_DIR / "05-documentazione" / "corpus-stem-overrides.json"

PLACEHOLDER_VALUES = {"", "-", "_", "—", "–"}
SCRIPT_VERSION = "1.0"


@dataclass(frozen=True)
class WorkbookSpec:
    filename: str
    prefix: str
    group: str
    sheet: str = "Risposte del modulo 1"


WORKBOOK_SPECS: tuple[WorkbookSpec, ...] = (
    WorkbookSpec(
        filename="Inseganti-attuali.xlsx",
        prefix="insegnanti_attuali",
        group="Insegnanti in servizio",
    ),
    WorkbookSpec(
        filename="Inseganti-futuri.xlsx",
        prefix="insegnanti_futuri",
        group="Insegnanti pre-service",
    ),
    WorkbookSpec(
        filename="Studenti.xlsx",
        prefix="studenti",
        group="Studenti",
    ),
)


@dataclass
class QAItem:
    question: str
    answer: str
    source_col: int


@dataclass
class Record:
    source_file: str
    source_sheet: str
    source_row: int
    group: str
    prefix: str
    record_code: str
    stem_base: str
    stem: str
    output_file: str
    item_count: int
    item_source_cols: list[int]


@dataclass
class BuildStats:
    skipped_blank_code: int = 0
    skipped_invalid_stem: int = 0
    skipped_empty_record: int = 0
    overrides_used: int = 0


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Build canonical corpus markdown files from dati/*.xlsx."
    )
    parser.add_argument(
        "--dati-dir",
        type=Path,
        default=DEFAULT_DATI_DIR,
        help=f"Directory with immutable source XLSX files (default: {DEFAULT_DATI_DIR}).",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=DEFAULT_OUTPUT_DIR,
        help=f"Destination directory for generated markdown files (default: {DEFAULT_OUTPUT_DIR}).",
    )
    parser.add_argument(
        "--manifest-out",
        type=Path,
        default=DEFAULT_MANIFEST_OUT,
        help=f"Manifest JSON path (written only with --write, default: {DEFAULT_MANIFEST_OUT}).",
    )
    parser.add_argument(
        "--overrides-file",
        type=Path,
        default=DEFAULT_OVERRIDES,
        help=f"Optional JSON map for legacy stem compatibility (default: {DEFAULT_OVERRIDES}).",
    )
    parser.add_argument(
        "--max-records",
        type=int,
        help="Optional global cap for generated records (debug only).",
    )
    parser.add_argument(
        "--write",
        action="store_true",
        help="Write markdown files and manifest. Without this flag the script is dry-run.",
    )
    parser.add_argument(
        "--overwrite",
        action="store_true",
        help="Allow overwriting existing markdown files in output-dir.",
    )
    return parser.parse_args()


def norm_space(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).replace("\xa0", " ")
    return re.sub(r"\s+", " ", text).strip()


def is_placeholder(value: Any) -> bool:
    text = norm_space(value)
    return text in PLACEHOLDER_VALUES


def stem_from_code(code: str) -> str:
    stem = norm_space(code).lower()
    stem = stem.replace("/", "-").replace("\\", "-")
    stem = re.sub(r"\s+", "-", stem)
    stem = re.sub(r"-{2,}", "-", stem).strip("-")
    return stem


def make_override_key(prefix: str, code: str) -> str:
    return f"{prefix}::{norm_space(code).lower()}"


def load_overrides(path: Path) -> dict[str, str]:
    if not path.exists():
        return {}
    payload = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(payload, dict):
        raise ValueError(f"Invalid overrides JSON at {path}: expected object at top-level.")
    normalized: dict[str, str] = {}
    for raw_key, raw_value in payload.items():
        key = norm_space(raw_key).lower()
        value = norm_space(raw_value)
        if not key or not value:
            continue
        normalized[key] = value
    return normalized


def extract_items(headers: tuple[Any, ...], row: tuple[Any, ...]) -> list[QAItem]:
    items: list[QAItem] = []
    for col in range(2, len(headers) + 1):
        question_raw = headers[col - 1] if col - 1 < len(headers) else None
        question = norm_space(question_raw)
        if not question:
            continue

        answer_raw = row[col - 1] if col - 1 < len(row) else None
        if is_placeholder(answer_raw):
            continue

        answer = norm_space(answer_raw)
        if not answer:
            continue

        items.append(QAItem(question=question, answer=answer, source_col=col))
    return items


def render_record_markdown(
    *,
    record_code: str,
    group: str,
    source_file: str,
    source_sheet: str,
    source_row: int,
    items: list[QAItem],
) -> str:
    lines: list[str] = []
    lines.append(f"# Risposta {record_code}")
    lines.append("")
    lines.append(f"**Gruppo di origine:** {group}")
    lines.append(
        f"**Sorgente dati:** dati/{source_file} · sheet: {source_sheet} · row: {source_row}"
    )
    lines.append(f"**Pipeline stage:** dati-to-corpus-v{SCRIPT_VERSION}")
    lines.append("")
    lines.append(f"- **Codice**: {record_code}")
    for item in items:
        lines.append(f"- **{item.question}**: {item.answer}")
    lines.append("")
    return "\n".join(lines)


def build_records(
    *,
    dati_dir: Path,
    output_dir: Path,
    overrides: dict[str, str],
    max_records: int | None,
) -> tuple[list[Record], dict[Path, str], BuildStats]:
    records: list[Record] = []
    file_contents: dict[Path, str] = {}
    stats = BuildStats()

    global_counter = 0

    for spec in WORKBOOK_SPECS:
        workbook_path = dati_dir / spec.filename
        if not workbook_path.exists():
            raise FileNotFoundError(f"Missing source workbook: {workbook_path}")

        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        if spec.sheet not in workbook.sheetnames:
            raise ValueError(
                f"Sheet '{spec.sheet}' not found in {workbook_path}. Available: {workbook.sheetnames}"
            )
        sheet = workbook[spec.sheet]

        rows = sheet.iter_rows(min_row=1, values_only=True)
        try:
            headers = next(rows)
        except StopIteration:
            workbook.close()
            continue

        stem_occurrences: Counter[str] = Counter()

        for source_row_idx, row in enumerate(rows, start=2):
            code = norm_space(row[0] if row else None)
            if not code:
                stats.skipped_blank_code += 1
                continue

            items = extract_items(headers, row)
            if not items:
                stats.skipped_empty_record += 1
                continue

            override_key = make_override_key(spec.prefix, code)
            if override_key in overrides:
                stem_base = overrides[override_key]
                stats.overrides_used += 1
            else:
                stem_base = stem_from_code(code)

            if not stem_base:
                stats.skipped_invalid_stem += 1
                continue

            stem_occurrences[stem_base] += 1
            occurrence = stem_occurrences[stem_base]
            stem = stem_base if occurrence == 1 else f"{stem_base}_{occurrence}"

            out_name = f"{spec.prefix}_{stem}.md"
            out_path = output_dir / out_name

            content = render_record_markdown(
                record_code=code,
                group=spec.group,
                source_file=spec.filename,
                source_sheet=spec.sheet,
                source_row=source_row_idx,
                items=items,
            )

            records.append(
                Record(
                    source_file=spec.filename,
                    source_sheet=spec.sheet,
                    source_row=source_row_idx,
                    group=spec.group,
                    prefix=spec.prefix,
                    record_code=code,
                    stem_base=stem_base,
                    stem=stem,
                    output_file=out_name,
                    item_count=len(items),
                    item_source_cols=[item.source_col for item in items],
                )
            )
            file_contents[out_path] = content

            global_counter += 1
            if max_records is not None and global_counter >= max_records:
                workbook.close()
                return records, file_contents, stats

        workbook.close()

    return records, file_contents, stats


def write_outputs(
    *,
    file_contents: dict[Path, str],
    overwrite: bool,
) -> None:
    for path in file_contents:
        if path.exists() and not overwrite:
            raise FileExistsError(
                f"Output file already exists: {path} (re-run with --overwrite to replace)."
            )
    for path, content in file_contents.items():
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(content, encoding="utf-8")


def build_manifest(
    *,
    args: argparse.Namespace,
    records: list[Record],
    stats: BuildStats,
    overrides: dict[str, str],
) -> dict[str, Any]:
    by_prefix = Counter(record.prefix for record in records)
    by_group = Counter(record.group for record in records)
    item_histogram = Counter(record.item_count for record in records)

    return {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "script": "scripts/build_corpus_from_dati.py",
        "script_version": SCRIPT_VERSION,
        "mode": "write" if args.write else "dry-run",
        "dati_dir": str(args.dati_dir),
        "output_dir": str(args.output_dir),
        "manifest_out": str(args.manifest_out),
        "overrides_file": str(args.overrides_file),
        "overrides_loaded": len(overrides),
        "records_total": len(records),
        "records_by_prefix": dict(sorted(by_prefix.items())),
        "records_by_group": dict(sorted(by_group.items())),
        "item_count_histogram": {str(k): v for k, v in sorted(item_histogram.items())},
        "skipped": asdict(stats),
        "records": [asdict(record) for record in records],
    }


def print_summary(manifest: dict[str, Any]) -> None:
    print("[build-corpus-from-dati] summary")
    print(f"  mode: {manifest['mode']}")
    print(f"  dati_dir: {manifest['dati_dir']}")
    print(f"  output_dir: {manifest['output_dir']}")
    print(f"  records_total: {manifest['records_total']}")
    print(f"  records_by_prefix: {manifest['records_by_prefix']}")
    print(f"  skipped: {manifest['skipped']}")


def main() -> int:
    args = parse_args()

    if not args.dati_dir.exists():
        raise SystemExit(f"dati directory does not exist: {args.dati_dir}")

    overrides = load_overrides(args.overrides_file)

    records, file_contents, stats = build_records(
        dati_dir=args.dati_dir,
        output_dir=args.output_dir,
        overrides=overrides,
        max_records=args.max_records,
    )

    manifest = build_manifest(args=args, records=records, stats=stats, overrides=overrides)
    print_summary(manifest)

    if not args.write:
        return 0

    write_outputs(file_contents=file_contents, overwrite=args.overwrite)
    args.manifest_out.parent.mkdir(parents=True, exist_ok=True)
    args.manifest_out.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(f"  wrote_files: {len(file_contents)}")
    print(f"  manifest: {args.manifest_out}")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except KeyboardInterrupt:
        raise SystemExit(130)
